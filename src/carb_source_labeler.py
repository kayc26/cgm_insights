"""Photo → carb-source classifier + decision-impact eval.

The one genuinely manual step in the v1 pipeline is identifying each meal's
carb source from its Bevel photo (87 meals, photo-reviewed by hand). Those
hand labels are the ground-truth eval set for automating that step with a
vision model.

GI is NOT re-estimated here: once the carb source is classified, GI comes
from the dataset's existing food→GI mapping, so label errors propagate to
recommendations exactly the way they would in production.

Eval metrics, in order of importance:
  1. flip rate — how many of the 87 meals get a *different recommendation*
                 when the classifier's label replaces the hand label.
                 A misclassification only matters if it changes the decision.
  2. accuracy  — raw carb-source classification accuracy, for diagnosis.
  3. confidence calibration — error rate of "low" vs "high" confidence
                 labels, the basis for a human-review routing rule.

Setup: export the study windows' photos from iPhone Photos into
data/photos_raw/ and run src/match_photos.py — it timestamp-matches them
into data/photos/<meal_id>_<k>.jpg. A meal aggregates food entries logged
within 15 minutes, so one meal may have several photos of different foods;
all of a meal's photos go to the classifier in a single call and are judged
as one meal. Then:

  python src/carb_source_labeler.py --label         # classify photos via API
  python src/carb_source_labeler.py                 # eval data/llm_carb_labels.csv

NOTE: data/llm_carb_labels.csv holds the recorded labeling pass that the
numbers in CASE_STUDY.md and DECISIONS.md are quoted from. Running --label
OVERWRITES it and the eval numbers will change with the model.

Human-in-the-loop correction (the product loop — auto-recognition proposes,
the user confirms or fixes; disagreements are often dominant-carb judgment
calls rather than vision errors, so the shipped metric is correction burden,
not raw accuracy):

  python src/carb_source_labeler.py --make-review   # review sheet (xlsx, dropdown)
  python src/carb_source_labeler.py --apply-review  # merge -> final labels CSV
"""

import argparse
import base64
import json
import sys
from pathlib import Path

import pandas as pd
from sklearn.linear_model import LinearRegression

MEALS = "data/meal_dataset.parquet"
LABELS = "data/llm_carb_labels.csv"
REVIEW = "data/review_carb_labels.xlsx"
FINAL = "data/final_carb_labels.csv"
PHOTOS = Path("data/photos")

# Cheapest vision-capable tier; the flip-rate eval is what qualifies (or
# disqualifies) it against a flagship model on this exact task.
MODEL = "claude-haiku-4-5"

PROMPT = """The photo(s) above show the foods logged as ONE meal (entries
within a 15-minute window). Identify the dominant carbohydrate source of the
meal as a whole.

Answer with exactly one label from this vocabulary (these are the categories
used throughout the dataset):

{vocab}

Rules:
- Pick the single food contributing the most carbohydrate across all photos.
- If no vocabulary label fits, use "other".
- confidence: "high" if the carb source is clearly visible, "medium" if
  partially obscured or inferred, "low" if you are mostly guessing.

Return JSON only: {{"carb_source": ..., "confidence": ...}}"""


def label_photos(vocab: list[str]) -> pd.DataFrame:
    import anthropic

    client = anthropic.Anthropic()

    by_meal: dict[int, list[Path]] = {}
    for photo in sorted(PHOTOS.glob("*.[jp][pn]g")):
        meal_id = int(photo.stem.split("_")[0])
        by_meal.setdefault(meal_id, []).append(photo)

    rows = []
    for meal_id, photos in sorted(by_meal.items()):
        image_blocks = [
            {
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": "image/png" if p.suffix == ".png" else "image/jpeg",
                    "data": base64.standard_b64encode(p.read_bytes()).decode(),
                },
            }
            for p in photos
        ]
        response = client.messages.create(
            model=MODEL,
            max_tokens=256,
            messages=[{
                "role": "user",
                "content": image_blocks
                + [{"type": "text", "text": PROMPT.format(vocab="\n".join(vocab))}],
            }],
            output_config={
                "format": {
                    "type": "json_schema",
                    "schema": {
                        "type": "object",
                        "properties": {
                            "carb_source": {"type": "string", "enum": vocab + ["other"]},
                            "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                        },
                        "required": ["carb_source", "confidence"],
                        "additionalProperties": False,
                    },
                }
            },
        )
        if response.stop_reason == "refusal":
            sys.exit(f"model refused on meal {meal_id}")
        result = json.loads(response.content[0].text)
        rows.append({"meal_id": meal_id, "llm_carb_source": result["carb_source"],
                     "confidence": result["confidence"]})
        print(f"meal {meal_id:3d} ({len(photos)} photo{'s' if len(photos) > 1 else ''})  "
              f"{result['carb_source']:30s} ({result['confidence']})")
    return pd.DataFrame(rows)


# ---- recommendation engine, mirrored verbatim from src/model.ipynb ----

FEATURES = ["GI", "DietaryCarbohydrates", "DietaryFatTotal", "DietaryFiber", "DietaryProtein", "pre_steps"]
MMOL_TO_MGDL = 18.0182
NO_SPIKE_MGDL = 100  # Lingo's published no-spike threshold


def recommend(model, meal_row):
    predicted_rise = model.predict(meal_row[FEATURES].to_frame().T)[0]
    peak_mgdl = (meal_row["baseline_glucose"] + predicted_rise) * MMOL_TO_MGDL
    if peak_mgdl < NO_SPIKE_MGDL:
        return "no action"

    variants = {}
    if meal_row["GI"] > 30:
        lower_gi = meal_row.copy()
        lower_gi["GI"] *= 0.6
        variants["gi swap"] = model.predict(lower_gi[FEATURES].to_frame().T)[0]
    veg_side = meal_row.copy()
    veg_side["DietaryFiber"] += 3
    veg_side["DietaryCarbohydrates"] += 4
    variants["veg side"] = model.predict(veg_side[FEATURES].to_frame().T)[0]
    return min(variants, key=variants.get)


def run_engine(meals: pd.DataFrame) -> pd.Series:
    """Full-pipeline counterfactual: retrain on the given GI column, then
    recommend — what the system would have done with these labels all along."""
    model = LinearRegression().fit(meals[FEATURES], meals["peak_rise"])
    return meals.apply(lambda row: recommend(model, row), axis=1)


def evaluate(meals: pd.DataFrame, labels: pd.DataFrame) -> None:
    gi_map = meals.drop_duplicates("carb_source").set_index("carb_source")["GI"]
    merged = meals.merge(labels, on="meal_id", how="inner", validate="1:1")
    if len(merged) < len(meals):
        print(f"note: only {len(merged)}/{len(meals)} meals have a photo label\n")

    correct = merged["llm_carb_source"] == merged["carb_source"]
    print(f"carb-source accuracy   {correct.sum()}/{len(merged)} ({correct.mean():.1%})")

    print("\nerror rate by confidence flag:")
    by_conf = merged.groupby("confidence").apply(
        lambda g: pd.Series({"n": len(g), "errors": (~(g["llm_carb_source"] == g["carb_source"])).sum()}),
        include_groups=False,
    )
    print(by_conf.to_string())

    # "other" and unseen labels have no GI in the mapping — treat as the
    # meal's hand GI so the flip metric isolates *wrong-label* impact, and
    # report them separately as would-be human-review cases.
    unmapped = ~merged["llm_carb_source"].isin(gi_map.index)
    if unmapped.any():
        print(f"\nunmapped labels routed to review: {unmapped.sum()}")
    llm_gi = merged["llm_carb_source"].map(gi_map).where(~unmapped, merged["GI"])

    baseline = run_engine(merged)
    swapped = run_engine(merged.assign(GI=llm_gi))
    flips = baseline != swapped
    print(f"\nrecommendation flips   {flips.sum()}/{len(merged)} meals ({flips.mean():.1%})")
    if flips.any():
        report = merged[flips][["meal_id", "carb_source", "llm_carb_source", "confidence"]].assign(
            was=baseline[flips], now=swapped[flips]
        )
        print(report.to_string(index=False))


def make_review(meals: pd.DataFrame) -> None:
    """Write the correction sheet: one row per photo-matched meal, the model's
    proposal prefilled in final_carb_source, a dropdown to fix it. The user
    only edits rows the model got wrong — that's the whole product loop."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    labels = pd.read_csv(LABELS)
    vocab = sorted(meals["carb_source"].unique()) + ["other"]
    with_photos = sorted({int(p.stem.split("_")[0]) for p in PHOTOS.glob("*.jpg")})

    sheet = meals[meals["meal_id"].isin(with_photos)][["meal_id", "meal_start"]].copy()
    sheet["meal_start"] = sheet["meal_start"].dt.tz_localize(None)  # Excel can't store tz
    sheet["photos"] = sheet["meal_id"].map(
        lambda mid: ", ".join(sorted(p.name for p in PHOTOS.glob(f"{mid}_*.jpg")))
    )
    sheet = sheet.merge(labels, on="meal_id", how="left")
    sheet["final_carb_source"] = sheet["llm_carb_source"]  # prefill; edit only the wrong ones
    sheet.to_excel(REVIEW, index=False)

    # dropdown on final_carb_source so corrections can't introduce typos
    wb = openpyxl.load_workbook(REVIEW)
    ws = wb.active
    vs = wb.create_sheet("vocab")
    for i, v in enumerate(vocab, start=1):
        vs.cell(row=i, column=1, value=v)
    vs.sheet_state = "hidden"
    dv = DataValidation(type="list", formula1=f"=vocab!$A$1:$A${len(vocab)}", allow_blank=True)
    ws.add_data_validation(dv)
    final_col = get_column_letter(sheet.columns.get_loc("final_carb_source") + 1)
    dv.add(f"{final_col}2:{final_col}{ws.max_row}")
    wb.save(REVIEW)

    pending = sheet["llm_carb_source"].isna().sum()
    print(f"wrote {REVIEW}: {len(sheet)} meals, {len(sheet) - pending} prefilled, "
          f"{pending} awaiting a prediction (run --label)")


def apply_review(meals: pd.DataFrame) -> None:
    """Merge the reviewed sheet into final labels and report the product
    metric: how much human correction the auto-labels actually needed."""
    sheet = pd.read_excel(REVIEW)
    sheet["final_carb_source"] = sheet["final_carb_source"].astype("string").str.strip()
    done = sheet[sheet["final_carb_source"].notna() & (sheet["final_carb_source"] != "")].copy()

    vocab = set(meals["carb_source"].unique()) | {"other"}
    bad = done[~done["final_carb_source"].isin(vocab)]
    if len(bad):
        sys.exit(f"labels outside the vocabulary (typo?): "
                 f"{bad[['meal_id', 'final_carb_source']].to_dict('records')}")

    auto = done["final_carb_source"] == done["llm_carb_source"]
    manual = done["llm_carb_source"].isna()
    corrected = ~auto & ~manual
    print(f"reviewed               {len(done)}/{len(sheet)} meals")
    print(f"auto-label accepted    {auto.sum()} ({auto.mean():.0%})")
    print(f"corrected              {corrected.sum()}")
    print(f"entered from scratch   {manual.sum()} (no prediction yet)")
    if corrected.any():
        print("\ncorrections by model confidence:")
        print(done[corrected].groupby("confidence").size().to_string())

    out = done[["meal_id", "final_carb_source"]].rename(columns={"final_carb_source": "carb_source"})
    out.to_csv(FINAL, index=False)
    print(f"\nwrote {FINAL}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--label", action="store_true", help="classify photos via the API")
    parser.add_argument("--make-review", action="store_true", help="write the manual-correction sheet")
    parser.add_argument("--apply-review", action="store_true", help="merge reviewed sheet into final labels")
    args = parser.parse_args()
    meals = pd.read_parquet(MEALS)
    if args.label:
        if not PHOTOS.exists() or not any(PHOTOS.iterdir()):
            sys.exit(f"no photos in {PHOTOS}/ — export originals to data/photos_raw/ "
                     "and run src/match_photos.py first")
        label_photos(sorted(meals["carb_source"].unique())).to_csv(LABELS, index=False)
        print(f"\nwrote {LABELS}")
    if args.make_review:
        make_review(meals)
    elif args.apply_review:
        apply_review(meals)
    else:
        evaluate(meals, pd.read_csv(LABELS))
